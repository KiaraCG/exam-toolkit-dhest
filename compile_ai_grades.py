#!/usr/bin/env python3
"""Compile configured exam scoring JSON into score and review XLSX/CSV files.

No AI calls, credentials or PDF processing. Run --check first. The shared
exam_config.json and rubric_manifest.json define all columns and score rules.
The grade CSV retains the historical two-header-row layout; XLSX adds totals,
part/problem summaries, explanations and source hashes. Review queues preserve
all exam-level flags; label mentions are navigation aids, not new item flags.
Existing outputs need explicit --overwrite and are backed up before replacement.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, DecimalException, localcontext
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from typing import Any, Iterable
import uuid
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from exam_common import (Config, ExamError, Item, load_config,
                         validate_result as common_validate, discover as common_discover)

VERSION = "EXAM_TOOLKIT_COMPILER_1.0.0"
OUTPUT_NAMES = ("ai_grades.xlsx", "human_reviews.xlsx", "ai_grades.csv", "human_reviews.csv")
EXCEL_ROW_LIMIT = 1_048_576
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def configure(config: Config) -> None:
    # One immutable configuration per CLI invocation; no hidden hard-coded exam.
    global CONFIG, ITEMS, BY_ID, PARTS, PART_ITEMS, PROBLEM_ITEMS, ITEM_INDEX
    global EXPECTED_COUNT, EXPECTED_TOTAL, PROBLEM_MAXIMA, BUNDLE_ID
    CONFIG = config
    ITEMS = config.items
    BY_ID = {i.label: i for i in ITEMS}
    PARTS = tuple(dict.fromkeys(i.part for i in ITEMS))
    PART_ITEMS = {p: tuple(i for i in ITEMS if i.part == p) for p in PARTS}
    PROBLEM_ITEMS = {p: tuple(i for i in ITEMS if i.problem == p)
                     for p in dict.fromkeys(i.problem for i in ITEMS)}
    PROBLEM_MAXIMA = {p: exact_sum(i.maximum for i in items) for p, items in PROBLEM_ITEMS.items()}
    ITEM_INDEX = {i.label: n for n, i in enumerate(ITEMS)}
    EXPECTED_COUNT, EXPECTED_TOTAL = len(ITEMS), config.total
    BUNDLE_ID = config.raw["exam_id"]


class InputError(ValueError):
    """An input is incomplete, inconsistent, or unsafe to export."""


@dataclass(frozen=True)
class ExamResult:
    path: Path
    number: str
    scores: dict[str, Decimal]
    explanations: dict[str, str]
    review: bool
    reason: str
    sha256: str

    @property
    def stem(self) -> str:
        return self.path.stem


@dataclass(frozen=True)
class Reference:
    scope: str
    part: str
    item: str
    literal: str
    basis: str


@dataclass(frozen=True)
class ReviewRow:
    exam: ExamResult
    reference: Reference


@dataclass(frozen=True)
class Formula:
    expression: str
    cached: Decimal


@dataclass
class Batch:
    results: list[ExamResult]
    warnings: list[str]
    pending: list[str]
    inventory_available: bool
    inventory_count: int


def strict_loads(text: str) -> Any:
    def unique(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result = {}
        for key, value in pairs:
            if key in result:
                raise InputError(f"duplicate JSON key: {key}")
            result[key] = value
        return result

    def invalid(value: str) -> Any:
        raise InputError(f"nonstandard JSON number: {value}")

    try:
        return json.loads(text, parse_float=Decimal, parse_int=Decimal,
                          object_pairs_hook=unique, parse_constant=invalid)
    except (ValueError, RecursionError, DecimalException) as exc:
        raise InputError(f"invalid JSON: {exc}") from exc


def number(value: Any, description: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise InputError(f"{description}: expected a number, not {value!r}")
    try:
        result = Decimal(str(value))
        as_float = float(result)
    except (ValueError, OverflowError) as exc:
        raise InputError(f"{description}: invalid numeric value") from exc
    if not result.is_finite() or not math.isfinite(as_float):
        raise InputError(f"{description}: number must be finite")
    if result != 0 and as_float == 0:
        raise InputError(f"{description}: number is too small for Excel")
    return result


def text_value(value: Any, description: str, allow_empty: bool = False) -> str:
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise InputError(f"{description}: expected {'a' if allow_empty else 'a nonempty'} string")
    if len(value.encode("utf-16-le", errors="surrogatepass")) // 2 > 32767:
        raise InputError(f"{description}: exceeds Excel's 32767-character cell limit; no truncation performed")
    if INVALID_XML.search(value):
        raise InputError(f"{description}: contains characters not representable in an XLSX cell")
    return value  # Preserve source wording and whitespace.


def exact_sum(values: Iterable[Decimal]) -> Decimal:
    values = list(values)
    # More than enough precision for the bounded input points, even when their
    # decimal encodings contain more digits than the usual half/quarter points.
    if not values:
        return Decimal(0)
    precision = max(40, max(len(v.as_tuple().digits) + abs(v.as_tuple().exponent) for v in values) + 10)
    with localcontext() as ctx:
        ctx.prec = precision
        return sum(values, Decimal(0))


def validate_manifest() -> None:
    if EXPECTED_COUNT + 2 > 16384:
        raise InputError("Too many criterion columns for XLSX; use a smaller exam scope")
    if len(ITEMS) != len(BY_ID) or exact_sum(i.maximum for i in ITEMS) != EXPECTED_TOTAL:
        raise InputError("Inconsistent shared manifest")


def read_result(path: Path) -> ExamResult:
    identifier, _ = CONFIG.identity(path.stem)
    raw = path.read_bytes()
    try:
        data = common_validate(strict_loads(raw.decode("utf-8-sig")), CONFIG, ordered=False)
    except (UnicodeError, ExamError) as exc:
        raise InputError(str(exc)) from exc
    scores = {g["rubric_item"]: number(g["awarded_points"], g["rubric_item"]) for g in data["grades"]}
    explanations = {g["rubric_item"]: text_value(g["explanation"], g["rubric_item"]) for g in data["grades"]}
    reason = text_value(data["human_review_reason"], "human_review_reason", allow_empty=True)
    return ExamResult(path, identifier, scores, explanations, data["human_review_recommended"],
                      reason, hashlib.sha256(raw).hexdigest())


def discover(folder: Path, extension: str) -> tuple[list[Path], list[str]]:
    try:
        return common_discover(CONFIG, folder, extension), []
    except ExamError as exc:
        return [], [str(exc)]


def load_batch(evals: Path, raw_exams: Path, require_complete: bool) -> Batch:
    if not evals.is_dir():
        raise InputError(f"missing evaluation directory: {evals}")
    files, errors = discover(evals, ".json")
    results = []
    for path in files:
        try:
            results.append(read_result(path))
        except (InputError, ExamError, OSError) as exc:
            errors.append(f"{path.name}: {exc}")
    if not files:
        errors.append(f"no recognized scoring JSON files in {evals}")
    warnings, pending = [], []
    available = raw_exams.is_dir()
    count = 0
    if available:
        pdfs, inventory_errors = discover(raw_exams, ".pdf")
        count = len(pdfs)
        # An inventory problem is not a fabricated zero score or another review flag.
        if require_complete:
            errors.extend(f"PDF inventory: {e}" for e in inventory_errors)
        else:
            warnings.extend(f"PDF inventory: {e}" for e in inventory_errors)
        if require_complete and not pdfs:
            errors.append("--require-complete: PDF inventory is empty")
        stems = {p.stem for p in pdfs}
        result_stems = {r.stem for r in results}
        pending = [p.name for p in pdfs if p.stem not in result_stems]
        if pending:
            message = "No valid matching JSON yet for: " + ", ".join(pending)
            (errors if require_complete else warnings).append(message)
        for result in results:
            if result.stem not in stems:
                message = f"{result.path.name}: no PDF with exactly matching filename stem in {raw_exams}"
                (errors if require_complete else warnings).append(message)
    else:
        message = f"PDF inventory unavailable at {raw_exams}; completeness cannot be checked"
        (errors if require_complete else warnings).append(message)
    if errors:
        raise InputError("Input checks failed; outputs left unchanged:\n  - " + "\n  - ".join(errors))
    return Batch(results, warnings, pending, available, count)


def extract_references(reason: str) -> tuple[Reference, ...]:
    """Literal manifest references only, not NLP or item-level review decisions.

    Whole-part aliases such as 4(c) are recognized when the manifest uses 4.c.
    Unrecognized shorthand/ranges stay in the preserved full reason. Longest
    matches first prevent a criterion mention becoming a whole-part flag.
    """
    hits, occupied = [], []
    def add(pattern, ref):
        for match in re.finditer(pattern, reason, re.IGNORECASE):
            if any(match.start() < b and a < match.end() for a, b in occupied):
                continue
            occupied.append(match.span())
            hits.append(Reference(ref.scope, ref.part, ref.item, match.group(), ref.basis))
    def literal(value):
        return r"(?<![\w.])" + re.escape(value) + r"(?![\w.]|\s*\.[A-Za-z0-9])"
    for item in sorted(ITEMS, key=lambda i: -len(i.label)):
        add(literal(item.label), Reference("criterion", item.part, item.label, "",
            "Criterion mentioned in exam-level reason; not a structured item flag"))
    for part in sorted(PARTS, key=len, reverse=True):
        add(literal(part), Reference("part", part, "", "",
            "Part mentioned in exam-level reason; not individual criterion flags"))
        if re.fullmatch(r"[0-9]+\.[A-Za-z0-9]+", part):
            first, last = part.split(".")
            add(literal(first + "(" + last + ")"), Reference("part", part, "", "",
                "Part mentioned in exam-level reason; not individual criterion flags"))
    for problem in PROBLEM_ITEMS:
        add(r"\b(?:problem|question|aufgabe)\s+" + re.escape(problem) + r"(?![\w.(])",
            Reference("problem", problem, "", "", "Problem mentioned in full reason; verify scope"))
    unique = {}
    for hit in hits:
        unique.setdefault((hit.scope, hit.part, hit.item), hit)
    def order(hit):
        if hit.item:
            return ITEM_INDEX[hit.item], 2
        group = PART_ITEMS.get(hit.part) if hit.scope == "part" else PROBLEM_ITEMS.get(hit.part)
        return ITEM_INDEX[group[0].label], 1 if hit.scope == "part" else 0
    return tuple(sorted(unique.values(), key=order))


def review_rows(results: list[ExamResult]) -> list[ReviewRow]:
    rows = []
    for exam in results:
        if not exam.review:
            continue
        refs = extract_references(exam.reason)
        if not refs:
            refs = (Reference("exam", "EXAM", "", "",
                              "No recognized explicit item/part reference; inspect the full reason and exam"),)
        rows.extend(ReviewRow(exam, ref) for ref in refs)
    return rows


def numeric(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def decimal_text(value: Decimal) -> str:
    if value == 0:
        return "0"
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def display_number(exam: ExamResult, width: int) -> str:
    return exam.number.zfill(width) if CONFIG.raw["identity_mode"] == "numeric" else exam.number


class WorkbookBuilder:
    """Create an XLSX with simple formulas and verified numeric cache values."""
    def __init__(self, title: str):
        from openpyxl import Workbook
        from openpyxl.workbook.properties import CalcProperties
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.properties.title = title
        self.wb.properties.creator = "Exam scoring compiler"
        self.wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        self.caches: dict[str, dict[str, Decimal]] = {}
        self.next_rows: dict[str, int] = {}

    def sheet(self, title: str, headers: list[str], widths: list[float], freeze: str = "B2"):
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = freeze
        ws.sheet_view.zoomScale = 90
        ws.print_options.horizontalCentered = False
        ws.print_title_rows = "1:1"
        self.row(ws, headers, imported=False)
        ws.row_dimensions[1].height = 32
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="173D4C")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for n, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(n)].width = width
        return ws

    def row(self, ws, values: Iterable[Any], imported: bool = True, caution: bool = False) -> int:
        from openpyxl.styles import Alignment, Font, PatternFill
        r = self.next_rows.get(ws.title, 1)
        self.next_rows[ws.title] = r + 1
        if r > EXCEL_ROW_LIMIT:
            raise InputError(f"{ws.title}: Excel row limit exceeded; use a smaller input batch")
        height = 22
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c)
            color = "17633B" if imported else "666666"
            if isinstance(value, Formula):
                cell.value = value.expression
                self.caches.setdefault(ws.title, {})[cell.coordinate] = value.cached
                color = "000000"
                cell.number_format = "General"
            elif isinstance(value, Decimal):
                cell.value = numeric(value)
                cell.number_format = "General"
            elif isinstance(value, str):
                text_value(value, f"{ws.title}!{cell.coordinate}", allow_empty=True)
                cell.value = value
                cell.data_type = "s"  # Untrusted source text never becomes a formula.
                cell.number_format = "@"
                if value.lstrip().startswith(("=", "+", "-", "@")):
                    cell.quotePrefix = True
                width = ws.column_dimensions[cell.column_letter].width or 13
                lines = sum(max(1, math.ceil(len(line) / max(8, width))) for line in value.splitlines())
                height = max(height, min(240, 15 * lines + 8))
            else:
                cell.value = value
            cell.font = Font(name="Calibri", size=11, color=color)
            cell.alignment = Alignment(horizontal="left" if isinstance(value, str) else "right",
                                       vertical="top", wrap_text=isinstance(value, str))
            if caution:
                cell.fill = PatternFill("solid", fgColor="FFF0D6")
            elif r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6F7")
        ws.row_dimensions[r].height = height
        return r

    def maxima(self, ws, values: list[Any]) -> None:
        from openpyxl.styles import Font, PatternFill
        r = self.row(ws, values, imported=False)
        for cell in ws[r]:
            cell.fill = PatternFill("solid", fgColor="DFECEB")
            cell.font = Font(name="Calibri", size=11, color="444444", bold=True)
        ws.print_title_rows = "1:2"

    def finish(self) -> None:
        from openpyxl.utils import get_column_letter
        for ws in self.wb.worksheets:
            if ws.title not in ("Scores", "Part scores"):
                ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0

    def save(self, path: Path) -> None:
        self.finish()
        self.wb.save(path)
        cache_formulas(path, self.wb.sheetnames, self.caches)
        # Headless readers must get totals, not empty cells. Do not resave the
        # loaded workbook: an ordinary openpyxl save would clear those caches.
        from openpyxl import load_workbook
        check = load_workbook(path, data_only=True, read_only=True)
        try:
            for title, cells in self.caches.items():
                # One streaming pass per sheet, not a random access per formula.
                pending = dict(cells)
                for row in check[title].iter_rows():
                    for cell in row:
                        if cell.coordinate not in pending:
                            continue
                        expected = pending.pop(cell.coordinate)
                        if (not isinstance(cell.value, (int, float)) or
                                not math.isclose(float(cell.value), float(expected), rel_tol=1e-13, abs_tol=1e-12)):
                            raise InputError(f"Formula cache verification failed at {title}!{cell.coordinate}")
                if pending:
                    raise InputError(f"Missing formula caches in {title}")
        finally:
            check.close()


def cache_formulas(path: Path, sheet_names: list[str], caches: dict[str, dict[str, Decimal]]) -> None:
    """Write caches only for formulas created by this compiler, never evaluate text.

    openpyxl stores formulas but does not evaluate them. Each generated SUM or
    link has a separately computed Decimal result. Preserve the formulas and
    insert their numeric <v> cache for viewers/headless readers. Excel is still
    instructed to recalculate normally when opened/edited.
    """
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    ns = {"m": main_ns}
    temporary = path.with_name(path.name + ".cache.tmp")
    try:
        with ZipFile(path) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as dest:
            for info in source.infolist():
                data = source.read(info.filename)
                match = re.fullmatch(r"xl/worksheets/sheet([0-9]+)\.xml", info.filename)
                if match:
                    title = sheet_names[int(match.group(1)) - 1]
                    known = caches.get(title, {})
                    root = ET.fromstring(data)
                    found = set()
                    for cell in root.findall(".//m:sheetData/m:row/m:c", ns):
                        if cell.find("m:f", ns) is None:
                            continue
                        address = cell.attrib["r"]
                        if address not in known:
                            raise InputError(f"Unexpected formula in {title}!{address}")
                        found.add(address)
                        cell.set("t", "n")
                        value = cell.find("m:v", ns)
                        if value is None:
                            value = ET.SubElement(cell, f"{{{main_ns}}}v")
                        value.text = decimal_text(known[address])
                    if found != set(known):
                        raise InputError(f"Missing generated formulas in {title}")
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dest.writestr(info, data)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def sum_formula(sheet: str, row: int, group: Iterable[Item], exam: ExamResult | None = None) -> Formula:
    from openpyxl.utils import get_column_letter
    group = tuple(group)
    indices = sorted(ITEM_INDEX[i.label] + 2 for i in group)
    runs = []
    for index in indices:
        if runs and index == runs[-1][-1] + 1:
            runs[-1].append(index)
        else:
            runs.append([index])
    refs = []
    for run in runs:
        left, right = get_column_letter(run[0]), get_column_letter(run[-1])
        refs.append(f"'{sheet}'!{left}{row}:{right}{row}")
    if len(refs) > 255:
        raise InputError("Too many noncontiguous groups for one Excel SUM; group manifest items by part")
    value = exact_sum(exam.scores[i.label] if exam else i.maximum for i in group)
    return Formula("=SUM(" + ",".join(refs) + ")", value)


def add_run_info(builder: WorkbookBuilder, batch: Batch, args, run_id: str, timestamp: str,
                 review_count: int) -> None:
    ws = builder.sheet("Run info", ["Field", "Value"], [29, 115], "B2")
    entries = [
        ("Compiler", VERSION), ("Bundle", BUNDLE_ID), ("Run ID", run_id), ("Generated UTC", timestamp),
        ("Input JSON directory", str(args.evals_dir.resolve())),
        ("PDF filename inventory", str(args.raw_exams_dir.resolve())),
        ("Valid scored exams", len(batch.results)), ("Rubric criteria", EXPECTED_COUNT),
        ("Maximum points", EXPECTED_TOTAL),
        ("Flagged exams", sum(exam.review for exam in batch.results)),
        ("Review reference rows", review_count),
        ("Pending PDF results", len(batch.pending) if batch.inventory_available else "Unknown; no PDF inventory"),
        ("Review authority", "Only the JSON's exam-level human_review_recommended flag selects an exam for review."),
        ("Reference rows", "Mentions in human_review_reason are navigation aids, not independently recorded item-level flags. Read the full reason; a mention may be context or a dependency."),
        ("Unlocalized issues", "Every flagged exam is in Flagged exams. With no recognized reference it also has an EXAM row in Review queue. General issues mixed with references remain in the full reason."),
        ("Reference syntax", "Recognizes exact configured criterion/part labels and simple parenthesized aliases; whole problem mentions require an explicit prefix. No range or natural-language scope inference. Read the full reason."),
        ("Text preservation", "Review reasons and item explanations retain the source text. XLSX strings are literal; CSV text starting with a formula trigger gets an apostrophe for safety."),
        ("Source linkage", "Explanations and Flagged exams identify source JSON filenames. Source SHA256 values below identify the bytes read by this compilation, not the PDFs/model used during scoring."),
        ("Provenance limit", "The generic scorer writes separate .provenance sidecars. This compiler records the JSON snapshot hashes but does not verify those sidecars or model conditions. Use the scorer --check for local input/provenance checks; a live resume also compares model settings. Neither checks correctness."),
        ("Identifiers", "Exam Number is the configured filename id group, not an inferred person identity. Original zero padding is retained; --width only adds display zeros. Full filename stems are retained in source columns."),
        ("Derived values", "XLSX totals/subtotals are formulas with initial numeric caches; CSV values are computed from the same scores. No new score rounding or half-point grid is imposed. Excel has finite numeric precision."),
        ("Snapshot", "This is a compilation snapshot, not a live dashboard or an automatic correction workflow. Editing a workbook never changes the scoring JSON or the other workbook/CSVs."),
        ("Completeness", "Only valid JSON files are included. Pending exams have no score row, not zero points. Use --require-complete for an exact filename-inventory gate."),
        ("Snapshot consistency", "Do not edit scoring JSON while compiling. Each JSON is read once; its hash is recorded. Finish scoring first for a stable complete batch."),
    ]
    for key, value in entries:
        builder.row(ws, [key, value], imported=False)
    for message in batch.warnings:
        builder.row(ws, ["Warning", message], imported=False, caution=True)
    for exam in batch.results:
        builder.row(ws, [str(exam.path), "SHA256 " + exam.sha256])


def write_workbooks(stage: Path, batch: Batch, args, rows: list[ReviewRow], run_id: str, timestamp: str) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    scores_book = WorkbookBuilder(CONFIG.raw["exam_title"] + " - scores")
    score_headers = ["Exam Number", *BY_ID, "Total"]
    scores = scores_book.sheet("Scores", score_headers, [18, *([10] * EXPECTED_COUNT), 12], "B3")
    last_score_column = get_column_letter(EXPECTED_COUNT + 1)
    scores_book.maxima(scores, ["Maximum Points", *(i.maximum for i in ITEMS),
                               Formula(f"=SUM(B2:{last_score_column}2)", EXPECTED_TOTAL)])
    for index, item in enumerate(ITEMS, 2):
        scores.cell(1, index).comment = Comment(f"{item.name}\nMaximum: {item.maximum}\nSource: configured manifest", "Exam compiler")
    for r, exam in enumerate(batch.results, 3):
        scores_book.row(scores, [display_number(exam, args.width), *(exam.scores[i.label] for i in ITEMS),
                                Formula(f"=SUM(B{r}:{last_score_column}{r})", exact_sum(exam.scores.values()))])
        scores.cell(r, 1).comment = Comment(f"Source: {exam.path}\nExam-level review flag: {exam.review}", "Exam compiler")
        if exam.review:
            scores.cell(r, 1).fill = PatternFill("solid", fgColor="FFF0D6")
        scores.cell(r, EXPECTED_COUNT + 2).font = Font(name="Calibri", size=11, bold=True, color="000000")
        scores.cell(r, EXPECTED_COUNT + 2).fill = PatternFill("solid", fgColor="DFECEB")

    summary = scores_book.sheet("Summary", ["Exam Number", "Exam File Stem", *("P" + p for p in PROBLEM_ITEMS), "Total", "Human review", "Full review reason", "Source JSON"],
                                [18, 23, *([9] * len(PROBLEM_ITEMS)), 12, 14, 90, 31])
    for r, exam in enumerate(batch.results, 3):
        scores_book.row(summary, [display_number(exam, args.width), exam.stem,
                                 *(sum_formula("Scores", r, PROBLEM_ITEMS[p], exam) for p in PROBLEM_MAXIMA),
                                 sum_formula("Scores", r, ITEMS, exam), exam.review, exam.reason, str(exam.path)],
                        caution=exam.review)

    parts = scores_book.sheet("Part scores", ["Exam Number", *PARTS, "Total"], [18, *([11] * len(PARTS)), 12], "B3")
    scores_book.maxima(parts, ["Maximum Points", *(sum_formula("Scores", 2, PART_ITEMS[p]) for p in PARTS), sum_formula("Scores", 2, ITEMS)])
    for r, exam in enumerate(batch.results, 3):
        scores_book.row(parts, [display_number(exam, args.width), *(sum_formula("Scores", r, PART_ITEMS[p], exam) for p in PARTS), sum_formula("Scores", r, ITEMS, exam)])

    explanation_headers = ["exam_number", "part_id", "rubric_item", "rubric_name", "max_points", "awarded_points", "explanation", "source_json", "exam_file"]
    explanations = scores_book.sheet("Explanations", explanation_headers, [18, 11, 15, 36, 14, 17, 95, 31, 23])
    for exam in batch.results:
        for item in ITEMS:
            scores_book.row(explanations, [display_number(exam, args.width), item.part, item.label, item.name,
                                          item.maximum, exam.scores[item.label], exam.explanations[item.label], str(exam.path), exam.stem])
    manifest = scores_book.sheet("Manifest", ["rubric_item", "part_id", "rubric_name", "max_points", "scoring_type"], [15, 12, 64, 14, 24])
    for item in ITEMS:
        scores_book.row(manifest, [item.label, item.part, item.name, item.maximum,
                                  (str(tuple(numeric(x) for x in item.allowed)) if item.mode == "allowed" else "Step " + str(item.step) if item.mode == "step" else "Continuous fractional")], imported=False)
    add_run_info(scores_book, batch, args, run_id, timestamp, len(rows))
    scores_book.save(stage / "ai_grades.xlsx")

    reviews_book = WorkbookBuilder(CONFIG.raw["exam_title"] + " - human review")
    headers = review_headers()
    queue = reviews_book.sheet("Review queue", headers, [18, 12, 90, 15, 14, 50, 24, 85, 31, 23])
    for row in rows:
        reviews_book.row(queue, review_values(row, args.width), caution=row.reference.scope == "exam")
    queue["F1"].comment = Comment("References are mentions, not structured item-level review flags. Always read the full explanation in column C.", "Exam compiler")
    flagged = reviews_book.sheet("Flagged exams", ["exam_number", "human_review_recommended", "human_review_reason", "recognized_references", "source_json", "exam_file"],
                                 [18, 30, 105, 45, 31, 23])
    context = reviews_book.sheet("All criteria", explanation_headers, [18, 11, 15, 36, 14, 17, 95, 31, 23])
    for exam in batch.results:
        if not exam.review:
            continue
        refs = extract_references(exam.reason)
        reviews_book.row(flagged, [display_number(exam, args.width), True, exam.reason,
                                  ", ".join(ref.item or ref.part for ref in refs) or "None recognized; exam-level review",
                                  str(exam.path), exam.stem], caution=not refs)
        for item in ITEMS:
            reviews_book.row(context, [display_number(exam, args.width), item.part, item.label, item.name,
                                      item.maximum, exam.scores[item.label], exam.explanations[item.label], str(exam.path), exam.stem])
    add_run_info(reviews_book, batch, args, run_id, timestamp, len(rows))
    # Important for a header-only queue: state that no exam has the source flag,
    # without inventing a data row that looks like a review item.
    if not rows:
        queue["A1"].comment = Comment("No valid input exam has human_review_recommended=true in this snapshot.", "Exam compiler")
    context["A1"].comment = Comment("ALL criteria of flagged exams are context, not separate flags for every criterion.", "Exam compiler")
    reviews_book.save(stage / "human_reviews.xlsx")


def review_headers() -> list[str]:
    # Retain the three historical columns first; add traceability/localization.
    return ["exam_number", "part_id", "explanation", "rubric_item", "reference_scope", "mapping_note", "matched_text", "item_explanation", "source_json", "exam_file"]


def review_values(row: ReviewRow, width: int) -> list[Any]:
    exam, ref = row.exam, row.reference
    return [display_number(exam, width), ref.part, exam.reason, ref.item, ref.scope, ref.basis,
            ref.literal, exam.explanations.get(ref.item, ""), str(exam.path), exam.stem]


def csv_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return decimal_text(value)
    if isinstance(value, str) and (value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n"))):
        return "'" + value
    return value


def write_csvs(stage: Path, results: list[ExamResult], rows: list[ReviewRow], width: int) -> None:
    import csv
    with (stage / "ai_grades.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        # Exact traditional matrix: no extra Total/review/identifier columns.
        writer.writerow(["Exam Number", *BY_ID])
        writer.writerow(["Maximum Points", *(decimal_text(i.maximum) for i in ITEMS)])
        for exam in results:
            writer.writerow([csv_safe(display_number(exam, width)), *(decimal_text(exam.scores[i.label]) for i in ITEMS)])
    with (stage / "human_reviews.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(review_headers())
        for row in rows:
            writer.writerow([csv_safe(value) for value in review_values(row, width)])


def existing_outputs(output_dir: Path) -> list[Path]:
    return [output_dir / name for name in OUTPUT_NAMES if os.path.lexists(output_dir / name)]


def publish(stage: Path, output_dir: Path, overwrite: bool, run_id: str) -> Path | None:
    """Stage all files first; back up explicit overwrites and restore on failure."""
    existing = existing_outputs(output_dir)
    if existing and not overwrite:
        raise InputError("Outputs already exist; use --overwrite or a new --output-dir: " + ", ".join(p.name for p in existing))
    if any(p.is_symlink() or not p.is_file() for p in existing):
        raise InputError("refusing to replace a symlink or non-file output")
    backup = None
    if existing:
        backup = output_dir / ".compile_backups" / run_id
        backup.mkdir(parents=True, exist_ok=False)
        for path in existing:
            shutil.copy2(path, backup / path.name)
    published: list[Path] = []
    try:
        for name in OUTPUT_NAMES:
            destination = output_dir / name
            if overwrite:
                os.replace(stage / name, destination)
            else:
                os.link(stage / name, destination)  # No-overwrite atomic publication.
            published.append(destination)
    except BaseException:
        for path in reversed(published):
            if backup and (backup / path.name).exists():
                shutil.copy2(backup / path.name, path)
            else:
                path.unlink(missing_ok=True)
        raise
    return backup


def parse_args(argv: list[str] | None = None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("exam_config.json"))
    parser.add_argument("--evals-dir", "--evals", dest="evals_dir", type=Path, default=None)
    parser.add_argument("--raw-exams-dir", type=Path, default=None, help="Filename inventory only; PDFs are never opened")
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--width", type=int, default=0, help="Optional minimum displayed exam-number width, e.g. 3; default retains exact suffix")
    parser.add_argument("--require-complete", action="store_true", help="Require an exact one-to-one RawExams filename/result inventory")
    parser.add_argument("--check", action="store_true", help="Read-only JSON/inventory validation; no outputs, credentials, or API calls")
    parser.add_argument("--overwrite", action="store_true", help="Replace generated outputs after backing up the previous versions")
    parser.add_argument("--version", action="version", version=VERSION)
    args = parser.parse_args(argv)
    if not 0 <= args.width <= 32:
        parser.error("--width must be between 0 and 32")
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    config = load_config(args.config)
    configure(config)
    args.evals_dir = args.evals_dir.resolve() if args.evals_dir else config.file("evals")
    args.raw_exams_dir = args.raw_exams_dir.resolve() if args.raw_exams_dir else config.file("exams")
    args.output_dir = args.output_dir.resolve() if args.output_dir else config.root
    validate_manifest()
    batch = load_batch(args.evals_dir, args.raw_exams_dir, args.require_complete)
    rows = review_rows(batch.results)
    print(f"{VERSION}: {len(batch.results)} valid exam(s), {EXPECTED_COUNT} criteria, {EXPECTED_TOTAL} points.")
    print(f"Review: {sum(exam.review for exam in batch.results)} flagged exam(s), {len(rows)} reference row(s).")
    for message in batch.warnings:
        print("WARNING: " + message, file=sys.stderr)
    print("References index mentions in the full exam-level reason, not structured item-level flags.")
    if args.check:
        print("Check passed. No files written; no API calls.")
        return 0
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise InputError('Install the dependency: python -m pip install "openpyxl>=3.1.5,<4"') from exc
    if args.output_dir.is_symlink() or (args.output_dir.exists() and not args.output_dir.is_dir()):
        raise InputError(f"output location is not a directory: {args.output_dir}")
    existing = existing_outputs(args.output_dir)
    if existing and not args.overwrite:
        raise InputError("Outputs already exist; use --overwrite or a new --output-dir: " + ", ".join(p.name for p in existing))
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_" + uuid.uuid4().hex[:10]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    # One compiler at a time in a destination; keep the lock file to avoid
    # inode races. Lock is advisory and does not lock files in Excel.
    try:
        import fcntl
    except ImportError as exc:
        raise InputError("This script uses the existing macOS/Linux local-filesystem workflow") from exc
    with (args.output_dir / ".compile_ai_grades.lock").open("a+") as lock:
        try:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise InputError("another compiler is writing this output directory") from exc
        try:
            with tempfile.TemporaryDirectory(prefix=".compile_", dir=args.output_dir) as temp:
                stage = Path(temp)
                write_workbooks(stage, batch, args, rows, run_id, timestamp)
                write_csvs(stage, batch.results, rows, args.width)
                for result in batch.results:
                    if hashlib.sha256(result.path.read_bytes()).hexdigest() != result.sha256:
                        raise InputError(f"{result.path.name} changed while compiling; outputs not published")
                backup = publish(stage, args.output_dir, args.overwrite, run_id)
        finally:
            fcntl.flock(lock.fileno(), fcntl.LOCK_UN)
    for name in OUTPUT_NAMES:
        print(f"Wrote {args.output_dir / name}")
    if backup:
        print(f"Previous outputs preserved in {backup}")
    print("Input JSON and PDFs unchanged. No AI requests were made.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (InputError, ExamError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
    except KeyboardInterrupt:
        print("\nInterrupted. Inputs are unchanged.", file=sys.stderr)
        raise SystemExit(130)
