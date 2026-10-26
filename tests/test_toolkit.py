"""Offline tests using invented submissions and temporary workspaces only."""
from __future__ import annotations
import base64
import contextlib
import copy
from dataclasses import replace
from decimal import Decimal
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import exam_common as common
import prepare_exam as prepare
import score_exams as scorer
import compile_ai_grades as compiler
from pdf_inputs import read_pdf, validate_rubric
from test_normalize_exam_pdfs import scan_pdf

ROOT = Path(__file__).resolve().parents[1]


def quiet(call, *args):
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        return call(*args)


class Workspace(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / 'exam'
        shutil.copytree(ROOT / 'examples/minimal_exam', self.root)
        self.path = self.root / 'exam_config.json'
        self.config = common.load_config(self.path)
        self.result = common.load_json(self.root / 'synthetic_evals/Exam001.json')
        self.review_result = common.load_json(self.root / 'synthetic_evals/Exam002.json')

    def save_config(self, raw):
        self.path.write_bytes(common.json_bytes(raw))
        self.config = common.load_config(self.path)
        return self.config

    def save_manifest(self, raw):
        self.config.file('manifest').write_bytes(common.json_bytes(raw))
        self.config = common.load_config(self.path)
        return self.config

    def evals(self):
        shutil.copytree(self.root / 'synthetic_evals', self.root / 'evals')
        return self.root / 'evals'


class ConfigTests(Workspace):
    def test_example_is_not_a_real_exam_profile(self):
        self.assertFalse(self.config.raw['approved_for_scoring'])
        self.assertEqual(self.config.total, 6)
        self.assertEqual(len(self.config.items), 5)

    def test_path_resolution_is_project_relative(self):
        self.assertEqual(self.config.file('exams'), self.root / 'RawExams')

    def test_unknown_configuration_key_is_not_ignored(self):
        raw=copy.deepcopy(self.config.raw); raw['expected_points_typo']=8
        with self.assertRaises(common.ExamError): self.save_config(raw)

    def test_path_escape_is_rejected(self):
        for value in ('../other', '/tmp/other'):
            with self.subTest(value=value):
                raw=copy.deepcopy(self.config.raw); raw['paths']['exams']=value
                with self.assertRaises(common.ExamError): self.save_config(raw)

    def test_symlink_escape_is_rejected(self):
        outside=Path(self.temp.name) / 'outside'; outside.mkdir()
        (self.root / 'link').symlink_to(outside, target_is_directory=True)
        raw=copy.deepcopy(self.config.raw); raw['paths']['exams']='link'
        with self.assertRaises(common.ExamError): self.save_config(raw)

    def test_distinct_paths(self):
        raw=copy.deepcopy(self.config.raw); raw['paths']['evals']='RawExams'
        with self.assertRaises(common.ExamError): self.save_config(raw)

    def test_total_checked(self):
        raw=copy.deepcopy(self.config.manifest); raw['expected_total_max_points']=53
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_bundle_id_checked(self):
        raw=copy.deepcopy(self.config.manifest); raw['exam_id']='OTHER'
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_duplicate_item_rejected(self):
        raw=copy.deepcopy(self.config.manifest); raw['items'][1]['rubric_item']='1.a.a'
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_unknown_dependency(self):
        raw=copy.deepcopy(self.config.manifest); raw['items'][1]['dependencies']=['missing']
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_dependency_cycle(self):
        raw=copy.deepcopy(self.config.manifest)
        raw['items'][2]['dependencies']=['2.b.b']
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_inconsistent_part_group(self):
        raw=copy.deepcopy(self.config.manifest); raw['items'][2]['problem_id']='POTHER'
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_mc_key_requires_text_and_option(self):
        raw=copy.deepcopy(self.config.manifest); del raw['items'][0]['correct_answer']
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_allowed_point_range(self):
        raw=copy.deepcopy(self.config.manifest); raw['items'][0]['allowed_points']=[0,1,2]
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_step_must_divide_maximum(self):
        raw=copy.deepcopy(self.config.manifest); raw['items'][1]['score_step']=0.3
        with self.assertRaises(common.ExamError): self.save_manifest(raw)

    def test_modes_are_not_universal_binary(self):
        raw=copy.deepcopy(self.config.manifest)
        raw['items'][0]['max_points']=4; raw['items'][0]['allowed_points']=[0,1,3,4]
        raw['expected_total_max_points']=9
        config=self.save_manifest(raw)
        self.assertEqual(config.items[0].validate_score(3),3)
        with self.assertRaises(common.ExamError): config.items[0].validate_score(2)

    def test_numeric_discovery_preserves_zeros(self):
        self.assertEqual([p.stem for p in common.discover(self.config,self.config.file('exams'),'.pdf')], ['Exam001','Exam002'])
        self.assertEqual(self.config.identity('Exam001'),('001',1))

    def test_numeric_duplicates_block(self):
        shutil.copy2(self.root/'RawExams/Exam001.pdf',self.root/'RawExams/1.pdf')
        with self.assertRaises(common.ExamError): common.discover(self.config,self.config.file('exams'),'.pdf')

    def test_unknown_filename_not_silently_ignored(self):
        (self.root/'RawExams/not-an-exam.pdf').write_bytes(b'%PDF-test')
        with self.assertRaises(common.ExamError): common.discover(self.config,self.config.file('exams'),'.pdf')

    def test_text_id_support(self):
        raw=copy.deepcopy(self.config.raw); raw['filename_pattern']=r'^(?P<id>[A-Za-z0-9-]+)$';raw['identity_mode']='text'
        config=self.save_config(raw)
        self.assertEqual(config.identity('candidate-X7'),('candidate-X7','candidate-x7'))


class ResultTests(Workspace):
    def test_good_result(self):
        self.assertEqual(common.validate_result(self.result,self.config),self.result)

    def test_no_missing_scores_as_zero(self):
        self.result['grades'].pop()
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_order_strict_for_scorer_flexible_for_compiler(self):
        self.result['grades'].reverse()
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)
        self.assertEqual(common.validate_result(self.result,self.config,ordered=False)['grades'][0]['rubric_item'],'1.a.a')

    def test_duplicate_criteria(self):
        self.result['grades'][1]=self.result['grades'][0]
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_wrong_name_or_maximum(self):
        for key,value in [('rubric_name','OTHER'),('max_points',99)]:
            with self.subTest(key=key):
                result=copy.deepcopy(self.result);result['grades'][0][key]=value
                with self.assertRaises(common.ExamError): common.validate_result(result,self.config)

    def test_bool_is_not_number(self):
        self.result['grades'][0]['awarded_points']=True
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_negative_and_excess_not_allowed(self):
        for value in [-0.1, 1.1]:
            self.result['grades'][0]['awarded_points']=value
            with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_fractional_mode_handling(self):
        self.result['grades'][2]['awarded_points']=0.25
        common.validate_result(self.result,self.config)
        self.result['grades'][1]['awarded_points']=0.25
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_binary_handling(self):
        self.result['grades'][0]['awarded_points']=0.5
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_explanation_required(self):
        self.result['grades'][0]['explanation']=' '
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_review_flag_reason_agreement(self):
        self.result['human_review_reason']='Inspect this'
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_no_extra_or_per_item_flag(self):
        self.result['grades'][0]['human_review_recommended']=True
        with self.assertRaises(common.ExamError): common.validate_result(self.result,self.config)

    def test_duplicate_keys_and_nan(self):
        for value in ['{"a":1,"a":2}', '{"a":NaN}', '{"a":Infinity}']:
            with self.subTest(value=value):
                with self.assertRaises(common.ExamError): common.strict_loads(value)

    def test_exact_decimal_sum(self):
        self.assertEqual(common.exact_sum([Decimal('0.1'),Decimal('0.2')]),Decimal('0.3'))

    def test_provider_schema_is_compact_and_local_rules_explicit(self):
        schema=common.output_schema(self.config)
        self.assertFalse(schema['additionalProperties'])
        self.assertFalse(schema['properties']['grades']['items']['additionalProperties'])
        self.assertEqual(len(schema['properties']['grades']['items']['properties']),5)


class PrepareAndPDFTests(Workspace):
    def test_rubric_headings_match(self):
        self.assertEqual(validate_rubric(self.config).page_count,1)

    def test_prompt_check_and_regeneration(self):
        quiet(prepare.main,['--config',str(self.path),'--check'])
        template=self.config.file('prompt_template')
        template.write_text(template.read_text()+'\nA new instructor instruction.\n')
        with self.assertRaises(common.ExamError):quiet(prepare.main,['--config',str(self.path),'--check'])
        with self.assertRaises(common.ExamError):quiet(prepare.main,['--config',str(self.path)])
        quiet(prepare.main,['--config',str(self.path),'--overwrite'])
        self.assertTrue(list((template.parent/'.prompt_backups').glob('*.txt')))

    def test_wrong_rubric_detected(self):
        shutil.copyfile(self.root/'RawExams/Exam001.pdf', self.config.file('rubric'))
        with self.assertRaises(common.ExamError):validate_rubric(self.config)

    def test_hash_approval_mode(self):
        raw=copy.deepcopy(self.config.raw)
        raw['rubric_validation']={'mode':'sha256','approved_sha256':common.sha256(self.config.file('rubric').read_bytes())}
        config=self.save_config(raw);validate_rubric(config)
        config.file('rubric').write_bytes(config.file('rubric').read_bytes()+b'\n')
        with self.assertRaises(common.ExamError):validate_rubric(config)

    def test_original_bad_scan_rejected_normalized_copy_accepted(self):
        from normalize_exam_pdfs import normalize_data
        p=self.root/'RawExams/Exam001.pdf';p.write_bytes(scan_pdf(image=(0,0,207,300)))
        with self.assertRaises(common.ExamError):read_pdf(p)
        data,_=normalize_data(p.read_bytes(),full_scan=True);p.write_bytes(data)
        self.assertEqual(read_pdf(p).simple_scan_pages,1)

    def test_rotation_rejected_without_rewrite(self):
        p=self.root/'RawExams/Exam001.pdf';data=scan_pdf(rotation=270);p.write_bytes(data)
        with self.assertRaises(common.ExamError):read_pdf(p)
        self.assertEqual(p.read_bytes(),data)

    def test_size_limit(self):
        with self.assertRaises(common.ExamError):read_pdf(self.root/'RawExams/Exam001.pdf',byte_limit=10)

    def test_check_no_credentials_or_files_written(self):
        before={p.relative_to(self.root):p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        with patch.object(scorer,'settings',side_effect=AssertionError('credentials accessed')):
            self.assertEqual(quiet(scorer.main,['--config',str(self.path),'--check']),0)
        after={p.relative_to(self.root):p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        self.assertEqual(before,after)

    def test_live_demo_blocked_before_credentials(self):
        with patch.object(scorer,'settings',side_effect=AssertionError('credentials accessed')):
            with self.assertRaisesRegex(common.ExamError,'approved_for_scoring'):
                quiet(scorer.main,['--config',str(self.path)])


class FakeClient:
    def __init__(self, result, status='completed'):
        self.created=[];self.deleted=[];self.requests=[]
        self.response=SimpleNamespace(status=status,output=[],output_text=json.dumps(result),id='response-synthetic',model='fake-model',usage=SimpleNamespace(input_tokens=12,output_tokens=20))
        self.files=SimpleNamespace(create=self.upload, delete=self.deleted.append)
        self.responses=SimpleNamespace(create=self.call)
    def upload(self, **kw):
        self.created.append(kw);return SimpleNamespace(id=f'file-synthetic-{len(self.created)}')
    def call(self, **kw):self.requests.append(kw);return self.response
    def __enter__(self):return self
    def __exit__(self,*args):pass


class ScorerTests(Workspace):
    def setUp(self):
        super().setUp()
        self.options=scorer.Settings('azure','https://example.invalid/openai/v1/','FAKE_MODEL','high',30000,900,2,'','TEST_ONLY_NOT_A_CREDENTIAL')
        self.rubric=validate_rubric(self.config)
        self.exam=read_pdf(self.root/'RawExams/Exam001.pdf')
        self.prompt=self.config.file('prompt').read_text()
        self.fingerprint=scorer.fingerprint(self.config,self.rubric,self.prompt,self.exam,self.options.public_record())
        self.out=self.root/'evals/Exam001.json'

    def test_azure_url_normalization(self):
        self.assertEqual(scorer.azure_base_url('https://example.invalid/'),'https://example.invalid/openai/v1/')
        with self.assertRaises(common.ExamError):scorer.azure_base_url('http://example.invalid')
        with self.assertRaises(common.ExamError):scorer.azure_base_url('https://user:secret@example.invalid')

    def test_secret_not_in_public_record_or_repr(self):
        self.assertNotIn(self.options.api_key,repr(self.options))
        self.assertNotIn(self.options.api_key,json.dumps(self.options.public_record()))

    def test_file_purpose_and_exact_bytes(self):
        client=FakeClient(self.result)
        scorer.upload_pdf(client,self.exam,self.options)
        self.assertEqual(client.created[0]['purpose'],'assistants')
        self.assertEqual(client.created[0]['file'][1],self.exam.data)
        scorer.upload_pdf(client,self.exam,replace(self.options,provider='openai'))
        self.assertEqual(client.created[1]['purpose'],'user_data')

    def test_openai_base_url(self):
        self.assertEqual(scorer.openai_base_url(''),'https://api.openai.com/v1/')
        self.assertEqual(scorer.openai_base_url('https://example.invalid/v1'),'https://example.invalid/v1/')
        with self.assertRaises(common.ExamError):scorer.openai_base_url('http://example.invalid/v1')
        with self.assertRaises(common.ExamError):scorer.openai_base_url('https://user:secret@example.invalid/v1')

    def test_inline_transport_embeds_exact_bytes_without_upload(self):
        client=FakeClient(self.result);options=replace(self.options,provider='openai',pdf_transport='inline')
        rubric,exam=scorer.upload_pdf(client,self.rubric,options),scorer.upload_pdf(client,self.exam,options)
        quiet(scorer.call_model,client,self.config,options,self.prompt,rubric,exam)
        scorer.delete_file(client,exam)
        parts=client.requests[0]['input'][0]['content']
        self.assertEqual(base64.b64decode(parts[1]['file_data'].split(',',1)[1]),self.exam.data)
        self.assertEqual(parts[0]['filename'],self.rubric.path.name);self.assertNotIn('file_id',parts[1])
        self.assertEqual((client.created,client.deleted),([],[]))

    def test_default_transport_keeps_provenance_record(self):
        self.assertNotIn('pdf_transport',self.options.public_record())
        self.assertEqual(replace(self.options,pdf_transport='inline').public_record()['pdf_transport'],'inline')

    def test_request_order_and_shape(self):
        client=FakeClient(self.result)
        result,meta=quiet(scorer.call_model,client,self.config,self.options,self.prompt,'RUBRIC','EXAM')
        request=client.requests[0];parts=request['input'][0]['content']
        self.assertEqual([parts[0]['file_id'],parts[1]['file_id']],['RUBRIC','EXAM'])
        self.assertNotIn('detail',parts[0])
        self.assertFalse(request['store']);self.assertEqual(request['truncation'],'disabled')
        self.assertEqual(result,self.result);self.assertEqual(meta['response_id'],'response-synthetic')

    def test_optional_detail_and_omitted_reasoning(self):
        client=FakeClient(self.result)
        quiet(scorer.call_model,client,self.config,replace(self.options,effort='',pdf_detail='high'),self.prompt,'R','E')
        self.assertNotIn('reasoning',client.requests[0])
        self.assertEqual(client.requests[0]['input'][0]['content'][0]['detail'],'high')

    def test_incomplete_response_never_accepted(self):
        with self.assertRaises(common.ExamError):quiet(scorer.call_model,FakeClient(self.result,'incomplete'),self.config,self.options,self.prompt,'R','E')
        self.assertFalse(self.out.exists())

    def test_refusal(self):
        client=FakeClient(self.result);client.response.output=[SimpleNamespace(content=[SimpleNamespace(type='refusal')])]
        with self.assertRaises(common.ExamError):quiet(scorer.call_model,client,self.config,self.options,self.prompt,'R','E')

    def test_invalid_json(self):
        client=FakeClient(self.result);client.response.output_text='not JSON'
        with self.assertRaises(common.ExamError):quiet(scorer.call_model,client,self.config,self.options,self.prompt,'R','E')

    def test_invalid_point_not_repaired(self):
        result=copy.deepcopy(self.result);result['grades'][0]['awarded_points']=9
        with self.assertRaises(common.ExamError):quiet(scorer.call_model,FakeClient(result),self.config,self.options,self.prompt,'R','E')

    def test_cleanup_failure_is_reported_not_fatal(self):
        client=FakeClient(self.result);client.files.delete=lambda _:(_ for _ in ()).throw(RuntimeError('fake failure'))
        quiet(scorer.delete_file,client,'file-id')

    def test_valid_resume_includes_review_flags(self):
        scorer.save_result(self.out,self.review_result,self.fingerprint,{})
        self.assertEqual(scorer.existing_status(self.out,self.config,self.fingerprint),'complete')

    def test_changed_inputs_or_model_block_resume(self):
        scorer.save_result(self.out,self.result,self.fingerprint,{})
        for key in ['exam_sha256','prompt_sha256','rubric_sha256','config_sha256','manifest_sha256','source_code_sha256','library_versions','request']:
            expected=copy.deepcopy(self.fingerprint);expected[key]='changed'
            with self.subTest(key=key):
                with self.assertRaises(common.ExamError):scorer.existing_status(self.out,self.config,expected)

    def test_old_results_without_sidecar_are_not_silently_reused(self):
        common.write_new(self.out,common.json_bytes(self.result))
        with self.assertRaisesRegex(common.ExamError,'no provenance'):scorer.existing_status(self.out,self.config,self.fingerprint)

    def test_modified_result_hash(self):
        scorer.save_result(self.out,self.result,self.fingerprint,{})
        self.out.write_bytes(self.out.read_bytes()+b'\n')
        with self.assertRaisesRegex(common.ExamError,'hash'):scorer.existing_status(self.out,self.config,self.fingerprint)

    def test_orphan_sidecar(self):
        common.write_new(scorer.provenance_path(self.out),b'{}')
        with self.assertRaisesRegex(common.ExamError,'Orphan'):scorer.existing_status(self.out,self.config,self.fingerprint)

    def test_interrupt_between_result_and_sidecar_is_detected(self):
        real=common.write_new
        def fail_sidecar(path,data):
            if path.parent.name=='.provenance':raise OSError('simulated interruption')
            real(path,data)
        with patch.object(scorer,'write_new',side_effect=fail_sidecar):
            with self.assertRaises(OSError):scorer.save_result(self.out,self.result,self.fingerprint,{})
        self.assertTrue(self.out.exists())
        with self.assertRaises(common.ExamError):scorer.existing_status(self.out,self.config,self.fingerprint)

    def test_no_overwrite(self):
        scorer.save_result(self.out,self.result,self.fingerprint,{})
        before=self.out.read_bytes()
        with self.assertRaises(FileExistsError):scorer.save_result(self.out,self.review_result,self.fingerprint,{})
        self.assertEqual(self.out.read_bytes(),before)

    def test_end_to_end_fake_provider_and_resume(self):
        raw=copy.deepcopy(self.config.raw);raw['approved_for_scoring']=True
        self.save_config(raw);quiet(prepare.main,['--config',str(self.path)])
        client=FakeClient(self.result)
        fake_module=SimpleNamespace(OpenAI=lambda **kwargs:client)
        with patch.dict(sys.modules,{'openai':fake_module}),patch.object(scorer,'settings',return_value=self.options):
            self.assertEqual(quiet(scorer.main,['--config',str(self.path),'--exam','Exam001.pdf']),0)
            self.assertEqual(quiet(scorer.main,['--config',str(self.path),'--exam','Exam001.pdf']),0)
        self.assertEqual(len(client.requests),1)
        self.assertEqual(len(client.created),2)
        self.assertEqual(set(client.deleted),{'file-synthetic-1','file-synthetic-2'})
        self.assertTrue(scorer.provenance_path(self.out).exists())

    def test_failed_response_cleans_up_both_uploaded_files(self):
        raw=copy.deepcopy(self.config.raw);raw['approved_for_scoring']=True
        self.save_config(raw)
        client=FakeClient(self.result,'incomplete')
        with patch.dict(sys.modules,{'openai':SimpleNamespace(OpenAI=lambda **kwargs:client)}),patch.object(scorer,'settings',return_value=self.options):
            self.assertEqual(quiet(scorer.main,['--config',str(self.path),'--exam','Exam001.pdf']),1)
        self.assertEqual(len(client.deleted),2)
        self.assertFalse(self.out.exists())


class CompilerTests(Workspace):
    def setUp(self):
        super().setUp();self.evals();compiler.configure(self.config)
        self.args=['--config',str(self.path),'--require-complete']

    def test_compile_csv_and_cached_totals(self):
        from openpyxl import load_workbook
        import csv
        quiet(compiler.main,self.args)
        with (self.root/'ai_grades.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertEqual(rows[0],['Exam Number',*[i.label for i in self.config.items]])
        self.assertEqual(rows[2][0],'001');self.assertEqual(rows[3][0],'002')
        self.assertEqual(len(rows[2]),6)
        wb=load_workbook(self.root/'ai_grades.xlsx',data_only=True)
        self.assertEqual(wb['Scores']['G2'].value,6)
        self.assertEqual(wb['Scores']['G3'].value,6)
        self.assertEqual(wb['Scores']['G4'].value,4)
        self.assertEqual(wb['Scores']['A3'].data_type,'s');wb.close()
        wb=load_workbook(self.root/'ai_grades.xlsx',data_only=False)
        self.assertEqual(wb['Scores']['G3'].value,'=SUM(B3:F3)');wb.close()

    def test_review_queue_preserves_authoritative_exam_reason(self):
        quiet(compiler.main,self.args)
        from openpyxl import load_workbook
        wb=load_workbook(self.root/'human_reviews.xlsx',data_only=True)
        self.assertEqual(wb['Flagged exams'].max_row,2)
        self.assertEqual(wb['Flagged exams']['C2'].value,self.review_result['human_review_reason'])
        self.assertEqual(wb['Review queue']['D2'].value,'1.a.a');wb.close()

    def test_unrecognized_reason_still_yields_exam_row(self):
        self.review_result['human_review_reason']='The scan is unreadable; page location uncertain.'
        (self.root/'evals/Exam002.json').write_bytes(common.json_bytes(self.review_result))
        batch=compiler.load_batch(self.root/'evals',self.root/'RawExams',True)
        rows=compiler.review_rows(batch.results)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0].reference.scope,'exam')

    def test_part_alias_and_criterion_distinct(self):
        refs=compiler.extract_references('Check 2(a) and 2.b.b; Problem 1 is context.')
        self.assertEqual({(r.scope,r.item or r.part) for r in refs},{('part','2.a'),('criterion','2.b.b'),('problem','1')})

    def test_low_score_does_not_create_review_flag(self):
        r=copy.deepcopy(self.result);r['grades'][0]['awarded_points']=0;r['grades'][0]['explanation']='Could be ambiguous; this is only a test.'
        (self.root/'evals/Exam001.json').write_bytes(common.json_bytes(r))
        (self.root/'evals/Exam002.json').unlink()
        batch=compiler.load_batch(self.root/'evals',self.root/'RawExams',False)
        self.assertEqual(compiler.review_rows(batch.results),[])

    def test_pending_results_not_zero_filled(self):
        (self.root/'evals/Exam002.json').unlink()
        batch=compiler.load_batch(self.root/'evals',self.root/'RawExams',False)
        self.assertEqual(len(batch.results),1);self.assertEqual(batch.pending,['Exam002.pdf'])
        with self.assertRaises(compiler.InputError):compiler.load_batch(self.root/'evals',self.root/'RawExams',True)

    def test_readonly_check_no_outputs(self):
        self.assertEqual(quiet(compiler.main,self.args+['--check']),0)
        self.assertFalse((self.root/'ai_grades.xlsx').exists())

    def test_overwrite_requires_flag_and_backs_up(self):
        quiet(compiler.main,self.args)
        before=(self.root/'ai_grades.xlsx').read_bytes()
        with self.assertRaises(compiler.InputError):quiet(compiler.main,self.args)
        quiet(compiler.main,self.args+['--overwrite'])
        backups=list((self.root/'.compile_backups').glob('*/ai_grades.xlsx'))
        self.assertEqual(len(backups),1);self.assertEqual(backups[0].read_bytes(),before)

    def test_invalid_input_does_not_overwrite_exports(self):
        quiet(compiler.main,self.args)
        before=(self.root/'ai_grades.xlsx').read_bytes()
        (self.root/'evals/Exam001.json').write_text('{}')
        with self.assertRaises((common.ExamError,compiler.InputError)):quiet(compiler.main,self.args+['--overwrite'])
        self.assertEqual(before,(self.root/'ai_grades.xlsx').read_bytes())

    def test_formula_like_text_remains_text(self):
        from openpyxl import load_workbook
        self.review_result['human_review_reason']='=DANGEROUS_EXTERNAL_FUNCTION()'
        (self.root/'evals/Exam002.json').write_bytes(common.json_bytes(self.review_result))
        quiet(compiler.main,self.args)
        wb=load_workbook(self.root/'human_reviews.xlsx');c=wb['Flagged exams']['C2']
        self.assertEqual(c.data_type,'s');self.assertTrue(c.quotePrefix);wb.close()
        self.assertIn("'=DANGEROUS",(self.root/'human_reviews.csv').read_text())

    def test_text_identifier_safe_for_csv(self):
        import csv
        raw=copy.deepcopy(self.config.raw);raw['filename_pattern']=r'^(?P<id>.+)$';raw['identity_mode']='text'
        self.save_config(raw)
        for folder,extension in [('RawExams','pdf'),('evals','json')]:
            (self.root/folder/f'Exam001.{extension}').rename(self.root/folder/f'=ABC.{extension}')
        quiet(compiler.main,self.args)
        with (self.root/'ai_grades.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertIn("'=ABC",[r[0] for r in rows])

    def test_fractional_scores_remain_unrounded(self):
        self.result['grades'][2]['awarded_points']=0.25
        (self.root/'evals/Exam001.json').write_bytes(common.json_bytes(self.result))
        quiet(compiler.main,self.args)
        from openpyxl import load_workbook
        wb=load_workbook(self.root/'ai_grades.xlsx',data_only=True)
        self.assertEqual(wb['Scores']['D3'].value,0.25);self.assertEqual(wb['Scores']['G3'].value,5.25);wb.close()

    def test_empty_review_file_still_has_headers(self):
        self.review_result['human_review_recommended']=False;self.review_result['human_review_reason']=''
        (self.root/'evals/Exam002.json').write_bytes(common.json_bytes(self.review_result))
        quiet(compiler.main,self.args)
        from openpyxl import load_workbook
        wb=load_workbook(self.root/'human_reviews.xlsx')
        self.assertEqual(wb['Flagged exams'].max_row,1);self.assertEqual(wb['Review queue'].max_row,1);wb.close()

    def test_noncontiguous_group_formula_sums_only_members(self):
        raw=copy.deepcopy(self.config.manifest);raw['items'][4]['problem_id']='1';raw['items'][4]['part_id']='1.b'
        self.save_manifest(raw)
        quiet(compiler.main,self.args)
        from openpyxl import load_workbook
        wb=load_workbook(self.root/'ai_grades.xlsx',data_only=True)
        # Summary: first two columns identify the exam; C is problem 1.
        self.assertEqual(wb['Summary']['C2'].value,3)
        self.assertEqual(wb['Summary']['D2'].value,3);wb.close()

    def test_input_json_unchanged(self):
        before={p.name:p.read_bytes() for p in (self.root/'evals').glob('*.json')}
        quiet(compiler.main,self.args)
        self.assertEqual(before,{p.name:p.read_bytes() for p in (self.root/'evals').glob('*.json')})


if __name__=='__main__':unittest.main()
